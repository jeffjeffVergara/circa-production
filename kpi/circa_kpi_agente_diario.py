#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
CIRCA — Agente KPI diario v3 (17-jul-2026)
Estructura de reporte basada en el Excel ejecutivo de Paola (CIRCA_KPI_*.xlsx)
+ definiciones oficiales v2 (GMV neto, hora Lima, pedidos con codigo Circa)
+ ventanas flexibles con comparaciones y NA.

ENV: SUPABASE_URL, SUPABASE_SERVICE_KEY, KPI_OUT_DIR (def .), ANTHROPIC_API_KEY (opc)
"""
import os, sys, json, datetime as dt
import pandas as pd
from supabase import create_client
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

AHORA = dt.datetime.now(dt.timezone.utc).replace(tzinfo=None) - dt.timedelta(hours=5)
HOY = AHORA.date()
INACTIVO_DIAS = 15
NORTE_GMV = 50_000_000  # S/50M — norte de GMV canalizado
METAS = {"afiliaciones_dia": 5, "pedidos_financiados_dia": 5, "recompra_7d": 65, "recompra_14d": 85}
METRICAS = ["afiliaciones", "pedidos", "pedidos_financiados", "bodegas_compraron",
            "bodegas_financiaron", "gmv", "monto_financiado", "monto_contado", "revenue_fee"]

sb = create_client(os.environ["SUPABASE_URL"], os.environ["SUPABASE_SERVICE_KEY"])

def fetch(vista, **eq):
    q = sb.table(vista).select("*")
    for k, v in eq.items(): q = q.eq(k, v)
    rows, page = [], 0
    while True:
        r = q.range(page*1000, page*1000+999).execute()
        rows += r.data
        if len(r.data) < 1000: return rows
        page += 1

def df_safe(vista):
    try: return pd.DataFrame(fetch(vista))
    except Exception as e:
        print(f"AVISO: {vista} fallo: {e}"); return pd.DataFrame()

# ============ 1. BASE DIARIA Y VENTANAS ============
df = pd.DataFrame(fetch("vw_kpi_diario_v2"))
if df.empty: sys.exit("vw_kpi_diario_v2 sin datos")
df["fecha"] = pd.to_datetime(df["fecha"]).dt.date
df = df.set_index("fecha").sort_index()
df = df.reindex(pd.date_range(df.index.min(), HOY, freq="D").date, fill_value=0)
df.index.name = "fecha"
INICIO = df.index.min()

PED_NIVEL = pd.DataFrame(sb.table("pedidos")
    .select("bodega_id,created_at,monto_financiado,numero,estado,bodegas!inner(es_test)")
    .eq("bodegas.es_test", False).not_.is_("numero","null")
    .in_("estado", ["confirmado","en_camino","pago_reportado","entregado","pagado","recibido","preventa_aceptada"])
    .limit(5000).execute().data)
if not PED_NIVEL.empty:
    PED_NIVEL["f"] = (pd.to_datetime(PED_NIVEL.created_at, utc=True, format="ISO8601") - pd.Timedelta(hours=5)).dt.date
    PED_NIVEL["fin"] = pd.to_numeric(PED_NIVEL.monto_financiado).fillna(0) > 0

def bodegas_dist(d0, d1, solo_fin=False):
    if PED_NIVEL.empty: return 0
    m = PED_NIVEL[(PED_NIVEL.f >= d0) & (PED_NIVEL.f <= d1)]
    if solo_fin: m = m[m.fin]
    return int(m.bodega_id.nunique())

def agg(d0, d1):
    if d1 < INICIO: return None
    sub = df.loc[max(d0, INICIO):d1]
    if sub.empty: return None
    out = sub[METRICAS].sum().round(2).to_dict()
    out["bodegas_compraron"] = bodegas_dist(max(d0, INICIO), d1)
    out["bodegas_financiaron"] = bodegas_dist(max(d0, INICIO), d1, solo_fin=True)
    return out

def pct(actual, previo, m):
    if actual is None or previo is None or not previo.get(m): return "NA"
    return round((actual[m]-previo[m])/previo[m]*100, 1)

ayer = agg(HOY-dt.timedelta(days=1), HOY-dt.timedelta(days=1)) or {m:0 for m in METRICAS}
u7  = agg(HOY-dt.timedelta(days=6), HOY)
u30 = agg(HOY-dt.timedelta(days=29), HOY)
mtd0 = HOY.replace(day=1); mtd = agg(mtd0, HOY)
pm = pd.Timestamp(mtd0) - pd.offsets.MonthBegin(1)
try: pf = pm.date().replace(day=HOY.day)
except ValueError: pf = (pd.Timestamp(mtd0)-dt.timedelta(days=1)).date()
mtd_prev = agg(pm.date(), min(pf, mtd0-dt.timedelta(days=1)))
mc_p = pd.Timestamp(HOY).to_period("M") - 1
mes_cerrado = agg(mc_p.start_time.date(), mc_p.end_time.date())
mes_cerrado_prev = agg((mc_p-1).start_time.date(), (mc_p-1).end_time.date())
acumulado = agg(INICIO, HOY)
lun = HOY - dt.timedelta(days=HOY.weekday())
sem_cur = agg(lun, HOY)
sem_w1  = agg(lun-dt.timedelta(days=7), lun-dt.timedelta(days=1))
sem_w4  = agg(lun-dt.timedelta(days=28), lun-dt.timedelta(days=22))

# Meses (cerrados + en curso) y semanas ISO completas
meses = []
m = pd.Timestamp(INICIO).to_period("M")
while m <= pd.Timestamp(HOY).to_period("M"):
    cur = agg(m.start_time.date(), min(m.end_time.date(), HOY))
    prev = agg((m-1).start_time.date(), (m-1).end_time.date())
    meses.append({"mes": str(m), "cerrado": m < pd.Timestamp(HOY).to_period("M"),
                  **(cur or {}), "gmv_vs_mes_ant_pct": pct(cur, prev, "gmv"),
                  "fin_vs_mes_ant_pct": pct(cur, prev, "pedidos_financiados")})
    m += 1
semanas = []
w = lun
while w >= INICIO - dt.timedelta(days=6):
    cur = agg(w, min(w+dt.timedelta(days=6), HOY))
    prev = agg(w-dt.timedelta(days=7), w-dt.timedelta(days=1))
    if cur: semanas.append({"semana_lunes": str(w), **cur,
        "gmv_vs_sem_ant_pct": pct(cur, prev, "gmv"),
        "fin_vs_sem_ant_pct": pct(cur, prev, "pedidos_financiados")})
    w -= dt.timedelta(days=7)

# ============ 2. VISTAS DE ESTADO ============
bodegas_kpi = df_safe("vw_kpi_bodega")
estado_uso  = df_safe("vw_bodega_estado_uso")
tiers       = df_safe("vw_tiers")
tiers_res   = df_safe("vw_tiers_resumen")
recompra    = df_safe("vw_recompra")
vs_hist     = df_safe("vw_kpi_bodega_vs_historico")
modelo_dist = df_safe("vw_modelo_distribuidor")
elegibles   = df_safe("vw_kpi_elegible_incremento")
pipeline    = df_safe("vw_kpi_pipeline_v2")
dq          = df_safe("vw_kpi_dq_v2")
comercial   = df_safe("vw_bodega_comercial")
ult_compra  = df_safe("vw_kpi_ultima_compra_v2")
kpi_ped     = df_safe("vw_kpi_pedidos")
bod_meta = pd.DataFrame(sb.table("bodegas").select("id,codigo_afiliado,contrato_firmado_at,es_test")
                        .eq("es_test", False).execute().data)

def join_com(d, key="bodega_id"):
    if d.empty or comercial.empty: return d
    return d.merge(comercial, on=key, how="left", suffixes=("", "_c"))

def join_cod(d, key="bodega_id"):
    if d.empty or bod_meta.empty: return d
    out = d.merge(bod_meta[["id","codigo_afiliado"]], left_on=key, right_on="id",
                  how="left", suffixes=("", "_b"))
    drop = [c for c in ("id_b","id") if c in out.columns and c != key]
    out = out.drop(columns=drop)
    cols = ["codigo_afiliado"] + [c for c in out.columns if c != "codigo_afiliado"]
    return out[cols]

bodegas_x = join_cod(join_com(bodegas_kpi))
estado_x  = join_cod(join_com(estado_uso))
tiers_x   = join_cod(join_com(tiers))
cobranza  = kpi_ped[(kpi_ped.get("vencido_activo")==True) & (kpi_ped.get("es_test")==False)] \
            if not kpi_ped.empty else pd.DataFrame()

# Stock: enroladas, activas 30d, inactivas
enroladas = int(bod_meta.contrato_firmado_at.notna().sum()) if not bod_meta.empty else 0
activas30 = 0
inactivos_hoy = []
if not ult_compra.empty:
    ult_compra["ultima_compra"] = pd.to_datetime(ult_compra["ultima_compra"]).dt.date
    uf = pd.to_datetime(ult_compra["ultimo_financiado"]).dt.date
    activas30 = int((uf >= HOY-dt.timedelta(days=30)).sum())
    inactivos_hoy = ult_compra[ult_compra["ultima_compra"] < HOY-dt.timedelta(days=INACTIVO_DIAS)]["bodega_id"].tolist()
prev = sb.table("kpi_snapshots_diario").select("payload").order("fecha", desc=True).limit(1).execute().data
prev_in = set((prev[0]["payload"].get("inactivos_ids") or []) if prev else [])
nuevos_inactivos = [i for i in inactivos_hoy if i not in prev_in]
tcount = dict(zip(tiers_res.tier, tiers_res.bodegas)) if not tiers_res.empty else {}
rec = recompra.iloc[0].to_dict() if not recompra.empty else {}

# Top vendedores enrolando (MTD)
top_enrol = pd.DataFrame()
if not bod_meta.empty and not comercial.empty:
    b = bod_meta.copy()
    b["f"] = (pd.to_datetime(b.contrato_firmado_at, utc=True, format="ISO8601") - pd.Timedelta(hours=5)).dt.date
    b = b[b.f.notna() & (b.f >= mtd0)]
    top_enrol = (b.merge(comercial, left_on="id", right_on="bodega_id")
                 .groupby(["vendedor","supervisor"], dropna=False).size()
                 .reset_index(name="enroladas_mtd").sort_values("enroladas_mtd", ascending=False))

# Ranking vendedores (compra) desde bodegas_x
vendedores = pd.DataFrame()
if not bodegas_x.empty and "vendedor" in bodegas_x.columns:
    vendedores = (bodegas_x.groupby(["vendedor","supervisor"], dropna=False)
        .agg(bodegas=("bodega","count"), pedidos_fin=("pedidos_financiados","sum"),
             gmv=("gmv_total","sum"), revenue=("revenue_circa","sum"))
        .round(2).reset_index().sort_values("gmv", ascending=False))
    if not top_enrol.empty:
        vendedores = vendedores.merge(top_enrol, on=["vendedor","supervisor"], how="outer").fillna({"enroladas_mtd":0,"bodegas":0,"pedidos_fin":0,"gmv":0,"revenue":0})

# ============ 2b. METRICAS OS (CEO/GROWTH/RIESGO) ============
import numpy as np
tel = pd.DataFrame(sb.table("bodegas").select("id,telefono_whatsapp").eq("es_test", False).execute().data)
def add_tel(d, key="bodega_id"):
    if d.empty or tel.empty: return d
    return d.merge(tel, left_on=key, right_on="id", how="left", suffixes=("","_t"))

# MAFM: bodegas que financiaron en el mes (actual vs anterior)
def mafm_mes(p):
    if PED_NIVEL.empty: return 0
    m = PED_NIVEL[PED_NIVEL.fin & (pd.to_datetime(PED_NIVEL.f).dt.to_period("M")==p)]
    return int(m.bodega_id.nunique())
p_act = pd.Timestamp(HOY).to_period("M")
mafm, mafm_prev = mafm_mes(p_act), mafm_mes(p_act-1)

# Cartera y PAR (sobre monto financiado de pedidos no pagados)
cart = kpi_ped[(kpi_ped.es_test==False) & (pd.to_numeric(kpi_ped.monto_financiado)>0) & (kpi_ped.pagado!=True)].copy() if not kpi_ped.empty else pd.DataFrame()
cartera = float(pd.to_numeric(cart.monto_financiado).sum()) if not cart.empty else 0.0
if not cart.empty:
    cart["da"] = pd.to_numeric(cart.dias_atraso).fillna(0)
    v7 = float(pd.to_numeric(cart[cart.da>7].monto_financiado).sum())
    v30 = float(pd.to_numeric(cart[cart.da>30].monto_financiado).sum())
    vencido_total = float(pd.to_numeric(cart[cart.vencido_activo==True].monto_financiado).sum())
    n_venc = int((cart.vencido_activo==True).sum())
else: v7=v30=vencido_total=0.0; n_venc=0
par7 = round(v7/cartera*100,1) if cartera else 0.0
par30 = round(v30/cartera*100,1) if cartera else 0.0

# TTFF (enrolamiento -> 1er financiamiento, mediana) y TBP (dias entre compras, mediana)
ttff = tbp = "NA"
if not PED_NIVEL.empty and not bod_meta.empty:
    fin1 = PED_NIVEL[PED_NIVEL.fin].groupby("bodega_id").f.min().reset_index(name="f1")
    bm = bod_meta.copy()
    bm["fe"] = (pd.to_datetime(bm.contrato_firmado_at, utc=True, format="ISO8601")-pd.Timedelta(hours=5)).dt.date
    j = fin1.merge(bm[["id","fe"]], left_on="bodega_id", right_on="id")
    j = j[j.fe.notna()]
    if len(j): ttff = round(float(np.median([(a-b).days for a,b in zip(j.f1,j.fe)])),1)
    difs = []
    for _,g in PED_NIVEL.sort_values("f").groupby("bodega_id"):
        fs = sorted(set(g.f))
        difs += [(b-a).days for a,b in zip(fs,fs[1:])]
    if difs: tbp = round(float(np.median(difs)),1)

primer_fin_mtd = 0
if not PED_NIVEL.empty:
    f1 = PED_NIVEL[PED_NIVEL.fin].groupby("bodega_id").f.min()
    primer_fin_mtd = int((f1 >= mtd0).sum())

# Utilizacion
util_prom = 0.0; sin_util = 0
if not bodegas_kpi.empty:
    bk = bodegas_kpi.copy()
    for c in ["linea_aprobada","linea_disponible"]: bk[c] = pd.to_numeric(bk[c]).fillna(0)
    con = bk[bk.linea_aprobada>0]
    if len(con): util_prom = round(float(((con.linea_aprobada-con.linea_disponible)/con.linea_aprobada).mean())*100,1)
    sin_util = int(((bk.linea_aprobada>0) & (pd.to_numeric(bk.dias_sin_pedir).fillna(999)>30)).sum())

# Health score v0 (25 frec/25 pago/20 util/20 recur/10 antig)
health = pd.DataFrame()
if not bodegas_x.empty:
    h = bodegas_x.copy()
    for c in ["pedidos_por_semana","pagos_puntuales","pedidos_pagados","linea_aprobada","linea_disponible","dias_sin_pedir","dias_enrolada"]:
        h[c] = pd.to_numeric(h.get(c)).fillna(0)
    h["s_frec"] = (h.pedidos_por_semana/1.0).clip(0,1)*25
    h["s_pago"] = np.where(h.pedidos_pagados>0, h.pagos_puntuales/h.pedidos_pagados, 0.5)*25
    h["s_util"] = np.where(h.linea_aprobada>0,(h.linea_aprobada-h.linea_disponible)/h.linea_aprobada,0)*20
    h["s_rec"]  = (1 - (h.dias_sin_pedir/30).clip(0,1))*20
    h["s_ant"]  = (h.dias_enrolada/90).clip(0,1)*10
    h["health"] = (h.s_frec+h.s_pago+h.s_util+h.s_rec+h.s_ant).round(0)
    health = h.sort_values("health", ascending=False)

# Vendedor quality score v0 (40 act/30 rec/20 cob/10 enrol)
rk_vend = pd.DataFrame()
if not bodegas_x.empty and "vendedor" in bodegas_x.columns:
    q = bodegas_x.copy()
    for c in ["pedidos_financiados","pedidos_pagados","pagos_puntuales"]: q[c]=pd.to_numeric(q[c]).fillna(0)
    q["financio"] = q.pedidos_financiados>0
    q["recompro"] = q.pedidos_financiados>1
    grp = q.groupby(["vendedor"], dropna=False).agg(afiliadas=("bodega","count"),
        pct_fin=("financio","mean"), pct_rec=("recompro","mean"),
        pagos=("pagos_puntuales","sum"), pagados=("pedidos_pagados","sum")).reset_index()
    grp["pct_cob"] = np.where(grp.pagados>0, grp.pagos/grp.pagados, 0.5)
    if not top_enrol.empty:
        grp = grp.merge(top_enrol.groupby("vendedor").enroladas_mtd.sum().reset_index(), on="vendedor", how="left")
    grp["enroladas_mtd"] = pd.to_numeric(grp.get("enroladas_mtd")).fillna(0)
    en_max = grp.enroladas_mtd.max() or 1
    grp["score"] = (40*grp.pct_fin + 30*grp.pct_rec + 20*grp.pct_cob + 10*grp.enroladas_mtd/en_max).round(0)
    rk_vend = grp.sort_values("score", ascending=False)

# ============ 3. SNAPSHOT ============
reporte = {"fecha": str(HOY), "inicio": str(INICIO), "acumulado": acumulado, "mtd": mtd,
    "mtd_vs_mes_ant": {m: pct(mtd, mtd_prev, m) for m in METRICAS},
    "mes_cerrado": mes_cerrado, "u7": u7, "u30": u30, "ayer": ayer,
    "sem_cur_vs_w1": {m: pct(sem_cur, sem_w1, m) for m in METRICAS},
    "sem_cur_vs_w4": {m: pct(sem_cur, sem_w4, m) for m in METRICAS},
    "enroladas": enroladas, "activas30": activas30, "tiers": tcount, "recompra": rec,
    "inactivos_total": len(inactivos_hoy), "inactivos_nuevos": len(nuevos_inactivos),
    "inactivos_ids": inactivos_hoy,
    "limbo": pipeline.to_dict("records") if not pipeline.empty else [],
    "dq": dq.to_dict("records") if not dq.empty else []}
sb.table("kpi_snapshots_diario").upsert({"fecha": str(HOY),
    "payload": json.loads(json.dumps(reporte, default=str))}).execute()

# ============ 4. ALERTAS / ANALISIS ============
alertas = []
if ayer["afiliaciones"] < METAS["afiliaciones_dia"]:
    alertas.append(f"Afiliaciones ayer {int(ayer['afiliaciones'])}/{METAS['afiliaciones_dia']}")
if ayer["pedidos_financiados"] < METAS["pedidos_financiados_dia"]:
    alertas.append(f"Pedidos financiados ayer {int(ayer['pedidos_financiados'])}/{METAS['pedidos_financiados_dia']}")
if nuevos_inactivos: alertas.append(f"{len(nuevos_inactivos)} bodegas cayeron a inactivas hoy")
if not dq.empty: alertas.append(f"{len(dq)} pedidos con montos inconsistentes (hoja DATA_QUALITY)")
viejos = pipeline[pipeline.dias_en_limbo > 7] if not pipeline.empty else pd.DataFrame()
if not viejos.empty:
    alertas.append(f"{len(viejos)} preventas >7d en limbo (S/{viejos.monto.astype(float).sum():.2f})")
if not cobranza.empty: alertas.append(f"{len(cobranza)} pedidos vencidos activos (hoja COBRANZA)")
analisis = ""
if os.environ.get("ANTHROPIC_API_KEY"):
    try:
        import anthropic
        msg = anthropic.Anthropic().messages.create(model="claude-sonnet-4-6", max_tokens=1200,
            messages=[{"role":"user","content":
                "Analista KPI de Circa (comercio+credito bodegas Peru, piloto). Analiza el JSON. "
                "Espanol peruano, directo, max 250 palabras: 1) lo mas importante (malo primero), "
                "2) tendencias (ignora NA), 3) alertas, 4) UNA accion concreta.\n\n"
                + json.dumps(reporte, default=str)}])
        analisis = msg.content[0].text
    except Exception as e: analisis = f"(analisis API fallo: {e})"

# ============ 5. EXCEL EJECUTIVO ============
F = lambda b=False, sz=10, col="000000": Font(name="Arial", bold=b, size=sz, color=col)
NAVY, GRAY, RED = "1F3864", "D9D9D9", "C00000"
fill = lambda c: PatternFill("solid", start_color=c)
thin = Border(*[Side(style="thin", color="BFBFBF")]*4)
MON, NUM, PCT_F = '"S/ "#,##0.00', '#,##0', '0.0"%"'

wb = Workbook(); ws = wb.active; ws.title = "DASHBOARD"
ws.sheet_view.showGridLines = False
for col, wdt in zip("ABCDEFGHI", [22,13,13,13,12,10,10,10,12]): ws.column_dimensions[col].width = wdt
r = 1
ws.cell(r,1,"CIRCA — DASHBOARD EJECUTIVO").font = F(True,16,NAVY); r+=1
ws.cell(r,1,f"Actualizado {AHORA.strftime('%d/%m/%Y %H:%M')} Lima · piloto ZOOM-DIMAX · GMV=neto, fechas Lima, pedidos con codigo Circa").font = F(False,9,"666666"); r+=2

def seccion(titulo):
    global r
    c = ws.cell(r,1,titulo); c.font = F(True,11,"FFFFFF"); c.fill = fill(NAVY)
    for i in range(2,10): ws.cell(r,i).fill = fill(NAVY)
    r += 1

def tabla(headers, rows, fmts):
    global r
    for j,h in enumerate(headers,1):
        c = ws.cell(r,j,h); c.font = F(True,9); c.fill = fill(GRAY); c.border = thin
    r += 1
    for row in rows:
        for j,(v,fm) in enumerate(zip(row,fmts),1):
            c = ws.cell(r,j,v); c.font = F(); c.border = thin
            if fm and not isinstance(v,str): c.number_format = fm
        r += 1
    r += 1

def fila_p(nombre, k):
    if k is None: return [nombre]+["NA"]*8
    t = round(k["gmv"]/k["pedidos"],2) if k["pedidos"] else 0
    return [nombre,k["gmv"],k["monto_financiado"],k["monto_contado"],k["revenue_fee"],
            int(k["pedidos"]),int(k["pedidos_financiados"]),int(k["bodegas_compraron"]),t]

seccion("RESULTADOS POR PERIODO")
tabla(["PERIODO","GMV","FINANCIADO","CONTADO","REVENUE","PEDIDOS","FINANC.","BODEGAS","TICKET"],
    [fila_p("ULTIMOS 7D",u7), fila_p(f"MTD {HOY.strftime('%b').upper()}",mtd),
     fila_p(f"MES CERRADO ({mc_p})",mes_cerrado), fila_p("ULTIMOS 30D",u30), fila_p("ACUMULADO",acumulado)],
    [None,MON,MON,MON,MON,NUM,NUM,NUM,MON])

seccion("COMPARACIONES (Δ% · NA = sin historia comparable)")
tabla(["COMPARACION","GMV","PED.FIN","AFILIACIONES","BODEGAS"],
    [["MTD vs mismo rango mes anterior", pct(mtd,mtd_prev,"gmv"), pct(mtd,mtd_prev,"pedidos_financiados"), pct(mtd,mtd_prev,"afiliaciones"), pct(mtd,mtd_prev,"bodegas_compraron")],
     ["Mes cerrado vs anterior", pct(mes_cerrado,mes_cerrado_prev,"gmv"), pct(mes_cerrado,mes_cerrado_prev,"pedidos_financiados"), pct(mes_cerrado,mes_cerrado_prev,"afiliaciones"), pct(mes_cerrado,mes_cerrado_prev,"bodegas_compraron")],
     ["Semana en curso vs anterior (W-1)", pct(sem_cur,sem_w1,"gmv"), pct(sem_cur,sem_w1,"pedidos_financiados"), pct(sem_cur,sem_w1,"afiliaciones"), pct(sem_cur,sem_w1,"bodegas_compraron")],
     ["Semana en curso vs comparable (W-4)", pct(sem_cur,sem_w4,"gmv"), pct(sem_cur,sem_w4,"pedidos_financiados"), pct(sem_cur,sem_w4,"afiliaciones"), pct(sem_cur,sem_w4,"bodegas_compraron")]],
    [None,PCT_F,PCT_F,PCT_F,PCT_F])

seccion("ADOPCION Y TIERS")
tabla(["ENROLADAS","ACTIVAS 30D","TIER 1 (<=7d)","TIER 2 (8-15d)","PROBO NO VOLVIO","ZOMBIE"],
    [[enroladas, f"{activas30} ({activas30/enroladas*100:.0f}%)" if enroladas else 0,
      int(tcount.get("tier1",0)), int(tcount.get("tier2",0)),
      int(tcount.get("tier3",tcount.get("probo_no_volvio",0))), int(tcount.get("zombie",0))]],
    [NUM,None,NUM,NUM,NUM,NUM])

seccion("METAS Y RIESGO")
tabla(["ENROL. AYER","PED.FIN AYER","RECOMPRA 7D","RECOMPRA 14D","INACTIVAS (+nuevas)","LIMBO","VENCIDOS","% NORTE S/50M"],
    [[f"{int(ayer['afiliaciones'])} / {METAS['afiliaciones_dia']}",
      f"{int(ayer['pedidos_financiados'])} / {METAS['pedidos_financiados_dia']}",
      f"{rec.get('pct_recompra_7d','NA')}% / {METAS['recompra_7d']}",
      f"{rec.get('pct_recompra_14d','NA')}% / {METAS['recompra_14d']}",
      f"{len(inactivos_hoy)} (+{len(nuevos_inactivos)})",
      f"{len(pipeline)} (S/{pipeline.monto.astype(float).sum():.0f})" if not pipeline.empty else "0",
      len(cobranza) if not cobranza.empty else 0,
      f"{(acumulado['gmv'] if acumulado else 0)/NORTE_GMV*100:.3f}%"]],
    [None]*8)

seccion("ALERTAS DEL DIA")
if alertas:
    for a in alertas:
        c = ws.cell(r,1,"• "+a); c.font = F(True,10,RED); r += 1
else:
    ws.cell(r,1,"Sin alertas contra metas").font = F(False,10,"2E7D32"); r += 1
r += 1
if analisis:
    seccion("ANALISIS")
    for line in analisis.split("\n"):
        ws.cell(r,1,line).font = F(False,9); r += 1

def hoja(nombre, d):
    if d is None or (hasattr(d,"empty") and d.empty): return
    w2 = wb.create_sheet(nombre)
    for row in dataframe_to_rows(d, index=False, header=True): w2.append(row)
    for c in w2[1]: c.font = F(True,9); c.fill = fill(GRAY)
    for col in w2.columns:
        w2.column_dimensions[col[0].column_letter].width = min(28, max(10, max((len(str(c.value or "")) for c in col[:50]))+2))

hoja("BODEGAS", bodegas_x)
hoja("VENDEDORES", vendedores)
hoja("ESTADO_USO", estado_x)
hoja("TIERS", tiers_x)
hoja("COBRANZA", cobranza)
hoja("RECOMPRA", recompra)
hoja("VS_HISTORICO", vs_hist)
hoja("MODELO_DIST", modelo_dist)
hoja("ELEGIBLES_LINEA", elegibles)
hoja("SEMANAS", pd.DataFrame(semanas))
hoja("MESES", pd.DataFrame(meses))
hoja("PIPELINE_LIMBO", pipeline)
hoja("DATA_QUALITY", dq)
hoja("DIARIO", df.reset_index())

out = os.environ.get("KPI_OUT_DIR", ".")
xlsx = f"{out}/circa_kpi_{HOY}.xlsx"; wb.save(xlsx)
md = f"{out}/circa_kpi_{HOY}.md"
with open(md,"w") as f:
    f.write(f"# Circa KPI — {HOY}\n\n## Alertas\n" + ("\n".join("- "+a for a in alertas) or "Sin alertas")
            + (f"\n\n## Analisis\n{analisis}" if analisis else "")
            + f"\n\n## Detalle\n```json\n{json.dumps(reporte, indent=2, default=str)}\n```\n")

# ============ 6. DASHBOARD HTML (Circa OS) ============
from dashboard_html import build as build_html
bx = add_tel(bodegas_x) if not bodegas_x.empty else bodegas_x
def _rows(d, cols, n=25):
    if d is None or d.empty: return []
    return d.head(n)[cols].fillna("").values.tolist()
usem = agg(lun-dt.timedelta(days=7), lun-dt.timedelta(days=1))
lh = ab = zb = []
if not bx.empty:
    bx["dsp"] = pd.to_numeric(bx.dias_sin_pedir).fillna(0)
    lhd = bx[(bx.dsp>=5)&(bx.dsp<=7)].sort_values("dsp")
    lh = _rows(lhd, ["bodega","telefono_whatsapp","vendedor","ultimo_pedido","dsp","linea_disponible"])
    abd = bx[bx.dsp>=15].copy()
    abd["banda"] = np.where(abd.dsp>=30,"30d+",np.where(abd.dsp>=20,"20d","15d"))
    ab = _rows(abd.sort_values("dsp",ascending=False), ["bodega","telefono_whatsapp","vendedor","ultimo_pedido","dsp","banda"])
zbd = add_tel(estado_x[estado_x.estado_uso=="enrolada_sin_pedido"]) if not estado_x.empty else pd.DataFrame()
zb = _rows(zbd, ["bodega","telefono_whatsapp","vendedor","dias_enrolada"]) if not zbd.empty else []
elx = elegibles.merge(health[["bodega_id","health"]], on="bodega_id", how="left") if not elegibles.empty and not health.empty else elegibles
vip = _rows(elx, ["bodega","linea_aprobada","linea_sugerida","pagos_puntuales","health"]) if not elx.empty else []
osd = {"fecha": f"{HOY} {AHORA.strftime('%H:%M')}", "alertas": alertas,
  "mafm": mafm, "mafm_delta": round((mafm-mafm_prev)/mafm_prev*100,1) if mafm_prev else "NA",
  "enrol_mtd": int(mtd["afiliaciones"]) if mtd else 0, "enrol_delta": pct(mtd, mtd_prev, "afiliaciones"),
  "enroladas": enroladas, "activas30": activas30, "primer_fin_mtd": primer_fin_mtd,
  "rec7": rec.get("pct_recompra_7d","NA"), "rec14": rec.get("pct_recompra_14d","NA"),
  "meta_rec7": METAS["recompra_7d"], "meta_rec14": METAS["recompra_14d"],
  "gmv_sem": (sem_cur or {}).get("gmv",0), "gmv_sem_delta": pct(sem_cur, usem, "gmv"),
  "gmv_mtd": (mtd or {}).get("gmv",0), "gmv_mtd_delta": pct(mtd, mtd_prev, "gmv"),
  "rev_mtd": (mtd or {}).get("revenue_fee",0),
  "ticket_mtd": round((mtd or {}).get("gmv",0)/max((mtd or {}).get("pedidos",1),1),0),
  "cartera": cartera, "par7": par7, "par30": par30, "vencido": vencido_total, "n_vencidos": n_venc,
  "util_prom": util_prom, "sin_util": sin_util, "ttff": ttff, "tbp": tbp,
  "llamar_hoy": lh, "zombies": zb, "abandono": ab, "vip": vip,
  "limbo": _rows(pipeline, ["numero","nombre_comercial","estado","monto","dias_en_limbo"]) if not pipeline.empty else [],
  "rk_vend": _rows(rk_vend.assign(pf=(rk_vend.pct_fin*100).round(0), pr=(rk_vend.pct_rec*100).round(0),
        pc=(rk_vend.pct_cob*100).round(0)), ["vendedor","score","afiliadas","pf","pr","pc","enroladas_mtd"]) if not rk_vend.empty else [],
  "subir": _rows(elegibles, ["bodega","linea_aprobada","linea_sugerida","pagos_puntuales"]) if not elegibles.empty else [],
  "cobrar": _rows(cobranza, ["numero","bodega","vendedor","monto_financiado","dias_atraso"]) if not cobranza.empty else [],
  "health": _rows(health.assign(fq=health.pedidos_por_semana.round(2)),
        ["bodega","health","fq","pagos_puntuales","s_util","dias_sin_pedir","segmento"], 40) if not health.empty else []}
html_path = f"{out}/circa_os_{HOY}.html"
open(html_path,"w").write(build_html(osd))
hoja("HEALTH_SCORE", health[["codigo_afiliado","bodega","health","s_frec","s_pago","s_util","s_rec","s_ant","segmento","vendedor"]] if not health.empty else pd.DataFrame())
hoja("RK_VENDEDOR_CALIDAD", rk_vend)
wb.save(xlsx)

try:
    for path, ct in [(xlsx,"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),(md,"text/markdown"),(html_path,"text/html")]:
        with open(path,"rb") as fh:
            sb.storage.from_("reportes-kpi").upload(path.split("/")[-1], fh.read(),
                file_options={"content-type": ct, "upsert": "true"})
    print("Reportes subidos a Storage/reportes-kpi")
except Exception as e:
    print(f"AVISO: upload Storage fallo ({e})")
print(f"OK — snapshot {HOY}. ALERTAS: {'; '.join(alertas) if alertas else 'ninguna'}")
