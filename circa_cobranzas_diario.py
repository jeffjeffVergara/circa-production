#!/usr/bin/env python3
"""
CIRCA — corte diario de cobranza.

Genera de una sola corrida:
  · CIRCA_Vencimientos_<fecha>.xlsx   (3 hojas: detalle, ventana, preventas sin entregar)
  · CIRCA_Correo_Vencimientos_<fecha>.md  (correo corto: vencido, lo que vence, responsables)

Uso:
    python3 circa_cobranzas_diario.py              # corte de hoy
    python3 circa_cobranzas_diario.py 2026-09-15   # corte de una fecha concreta

Solo LECTURA sobre la base. No escribe nada en Supabase.

Notas de implementación (ver CLAUDE.md):
  · fetch_all() pagina SIEMPRE: el tope de 1000 filas es del servidor PostgREST,
    .limit() no sirve.
  · Las fechas se convierten a hora Perú con astimezone antes de comparar. La base
    guarda UTC y comparar en UTC corre el día.
  · parse_ts normaliza la fracción de segundo a 6 dígitos: en Python < 3.11
    fromisoformat revienta con 4 o 5 dígitos y un except silencioso lo convierte
    en "no pagó".
  · Los montos salen de app/services/fees.total_pagar_desde_pedido, la misma
    función del back office y del WhatsApp del grupo.
"""
from __future__ import annotations

import os
import sys
from datetime import date, datetime, timedelta
from zoneinfo import ZoneInfo

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

REPO = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, REPO)

from dotenv import load_dotenv  # noqa: E402

load_dotenv(os.path.join(REPO, ".env"))
from supabase import create_client  # noqa: E402

from app.services.fees import hoy_peru, total_pagar_desde_pedido  # noqa: E402

TZ = ZoneInfo("America/Lima")
UTC = ZoneInfo("UTC")
DIMAX = "d1a2b3c4-0001-4000-8000-000000000002"

DIAS_ES = ["lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo"]
DIAS_AB = ["lun", "mar", "mié", "jue", "vie", "sáb", "dom"]
MES_ES = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "set", "oct", "nov", "dic"]


def fmt_dia(d: date) -> str:
    return f"{DIAS_ES[d.weekday()]} {d.day}"


def fmt_dia_ab(d: date) -> str:
    return f"{DIAS_AB[d.weekday()]} {d.day}"


def fmt_fecha(d: date) -> str:
    return f"{d.day}-{MES_ES[d.month - 1]}-{d.year}"


HOY = date.fromisoformat(sys.argv[1]) if len(sys.argv) > 1 else hoy_peru()

_url = os.environ["SUPABASE_URL"]
_key = (os.environ.get("SUPABASE_SERVICE_KEY") or os.environ.get("SUPABASE_KEY")
        or os.environ.get("SUPABASE_ANON_KEY"))
cli = create_client(_url, _key)


def fetch_all(table: str, select: str, filters=None) -> list[dict]:
    """Pagina con .range(). NUNCA usar .execute() pelado ni .limit()."""
    out, off = [], 0
    while True:
        q = cli.table(table).select(select)
        for f in filters or []:
            q = getattr(q, f[0])(*f[1:])
        data = q.range(off, off + 999).execute().data or []
        out += data
        if len(data) < 1000:
            return out
        off += 1000


def parse_ts(ts):
    """timestamptz -> datetime aware. Normaliza la fracción a 6 dígitos (Python 3.10)."""
    if not ts:
        return None
    s = str(ts).replace("Z", "+00:00")
    if "." in s:
        head, _, tail = s.partition(".")
        digits = ""
        for ch in tail:
            if ch.isdigit():
                digits += ch
            else:
                tail = tail[len(digits):]
                break
        else:
            tail = ""
        s = f"{head}.{digits[:6].ljust(6, '0')}{tail}"
    try:
        dt = datetime.fromisoformat(s)
    except (ValueError, TypeError):
        return None
    return dt.replace(tzinfo=UTC) if dt.tzinfo is None else dt


def dia_peru(ts):
    dt = parse_ts(ts)
    return dt.astimezone(TZ).date() if dt else None


def corto(nombre: str | None) -> str:
    """'ESTRADA TRAVEZAN CARLOS ALBERTO' -> 'Estrada Carlos' (regla del agente diario)."""
    p = (nombre or "Sin vendedor").split()
    if len(p) >= 3:
        return f"{p[0].title()} {p[2].title()}"
    return " ".join(x.title() for x in p)


def solo_apellido(etiqueta: str) -> str:
    """'V0014 Peña Shirley' -> 'V0014 Peña' (formato de las tablas del correo)."""
    p = etiqueta.split()
    return " ".join(p[:2]) if len(p) >= 2 else etiqueta


# bodega_vendedores guarda los días en minúsculas SIN TILDE ('miercoles');
# historico_zoom los guarda en MAYÚSCULAS. Acá solo se presentan bonito.
DIAS_BONITO = {
    "lunes": "Lunes", "martes": "Martes", "miercoles": "Miércoles",
    "miércoles": "Miércoles", "jueves": "Jueves", "viernes": "Viernes",
    "sabado": "Sábado", "sábado": "Sábado", "domingo": "Domingo",
}


def dia_visita_bonito(d: str | None) -> str:
    """'miercoles' -> 'Miércoles'. Sin ruta -> aviso explícito, nunca celda vacía."""
    if not d:
        return "— sin ruta"
    k = str(d).strip().lower()
    return DIAS_BONITO.get(k, str(d).strip().title())


# ─────────────────────────── datos ───────────────────────────

PED_COLS = ("id,numero,estado,monto_financiado,fee_monto,monto_total_credito,total,plazo_dias,"
            "fecha_vencimiento,fecha_entregado,entregado_at,fecha_pagado,pagado_at,"
            "pago_cliente_sustento_subido_at,metodo_pago,updated_at,bodega_id,"
            "preventa_aceptada_at")

pedidos = fetch_all("pedidos", PED_COLS, [("eq", "distribuidor_id", DIMAX)])

# ── Abonos (pagos parciales) ───────────────────────────────────────────────
# OJO: total_pagar_desde_pedido() sabe descontar abonos, pero los busca con
# _abonos_de_pedido(), que usa app.services.db.sb — un cliente que SOLO existe
# dentro de la app Flask. Desde un script suelto revienta con "supabase_key is
# required" y su `except Exception: return []` traga el error y responde "sin
# abonos". El reporte entonces cobra el total completo sin avisar.
# Por eso los traemos con NUESTRO cliente y se los pasamos explícitos.
# [Confirmado 22-sep-2026: RABANAL CRC-205 salía S/106 debiendo S/56, y
#  AVILA CRC-240 salía S/51.50 debiendo S/1.]
abonos_por_pedido: dict[str, list[dict]] = {}
for a in fetch_all("abonos", "pedido_id,monto,fecha_pago,anulado",
                   [("eq", "anulado", False)]):
    abonos_por_pedido.setdefault(a["pedido_id"], []).append(
        {"monto": a["monto"], "fecha_pago": a["fecha_pago"]})
bodegas = {b["id"]: b for b in fetch_all("bodegas", "id,razon_social,telefono_whatsapp,es_test")}
vendedores = {v["id"]: v for v in fetch_all("vendedores", "id,codigo,nombre,supervisor")}
mapa: dict[str, str] = {}
# ── Día de visita ──────────────────────────────────────────────────────────
# OJO [trampa ya pisada 3 veces, ver CLAUDE.md]: dia_visita TIENE que estar en
# este select. Si no está, llega vacío y la columna sale en blanco SIN ERROR —
# un fallo técnico disfrazado de "esta bodega no tiene ruta".
# Formato en bodega_vendedores: minúsculas sin tilde ('miercoles'). En
# historico_zoom es MAYÚSCULAS; acá usamos bodega_vendedores, igual que el
# resto del script, y lo mostramos capitalizado.
ruta_visita: dict[str, str] = {}
for m in fetch_all("bodega_vendedores", "bodega_id,vendedor_id,activo,dia_visita",
                   [("eq", "activo", True)]):
    mapa.setdefault(m["bodega_id"], m["vendedor_id"])
    if m["bodega_id"] not in ruta_visita:
        ruta_visita[m["bodega_id"]] = m.get("dia_visita")

cartera, preventas = [], []
for p in pedidos:
    b = bodegas.get(p["bodega_id"]) or {}
    if b.get("es_test") or str(p.get("numero") or "").startswith("TEST"):
        continue
    mf = float(p.get("monto_financiado") or 0)
    if mf <= 0:
        continue
    if (p.get("metodo_pago") or "").lower() == "efectivo" or p.get("estado") == "pagado_cash":
        continue  # cash nunca genera cobranza financiada

    v = vendedores.get(mapa.get(p["bodega_id"])) or {}
    entrega = dia_peru(p.get("fecha_entregado") or p.get("entregado_at"))
    abonos = abonos_por_pedido.get(p.get("id")) or []
    fila = {
        "numero": p["numero"], "estado": p["estado"], "mf": mf,
        "fee": float(p.get("fee_monto") or 0), "plazo": int(p.get("plazo_dias") or 7),
        "bodega": b.get("razon_social"), "cel": b.get("telefono_whatsapp"),
        "vendedor": f"{v.get('codigo', '')} {corto(v.get('nombre'))}".strip() or "Sin vendedor",
        "dia_visita": dia_visita_bonito(ruta_visita.get(p["bodega_id"])),
        "supervisor": v.get("supervisor"), "entrega": entrega,
        "vence": date.fromisoformat(p["fecha_vencimiento"]) if p.get("fecha_vencimiento") else None,
        "abonado": sum(float(a["monto"] or 0) for a in abonos),
        "n_abonos": len(abonos),
    }

    if p["estado"] == "entregado":
        tot = total_pagar_desde_pedido(p, 0, HOY, abonos=abonos or None)
        fila["total"] = tot["total_pagar"]
        fila["mora"] = tot.get("mora_monto", 0)
        fila["dias"] = (fila["vence"] - HOY).days if fila["vence"] else None
        fila["mora_desde"] = entrega + timedelta(days=32) if entrega else None
        cartera.append(fila)
    elif p["estado"] in ("confirmado", "preventa_aceptada"):
        fila["total"] = mf + fila["fee"]
        preventas.append(fila)

cartera.sort(key=lambda x: (x["vence"] or date.max, -x["total"]))
preventas.sort(key=lambda x: -x["mf"])

vencidos = [x for x in cartera if x["dias"] is not None and x["dias"] < 0]
ventana = [x for x in cartera if x["dias"] in (0, 1, 2)]
TOT_CART = sum(x["total"] for x in cartera)
TOT_VENC = sum(x["total"] for x in vencidos)
TOT_VENT = sum(x["total"] for x in ventana)


def agrupar(clave):
    g: dict[str, list] = {}
    for x in cartera:
        k = x[clave] or "Sin asignar"
        r = g.setdefault(k, [0, 0, 0.0, 0.0])
        r[0] += 1
        r[3] += x["total"]
        if x["dias"] is not None and x["dias"] < 0:
            r[1] += 1
            r[2] += x["total"]
    return dict(sorted(g.items(), key=lambda kv: (-kv[1][2], -kv[1][3])))


por_sup = agrupar("supervisor")
por_vend = agrupar("vendedor")

# ─────────────────────────── Excel ───────────────────────────

AZUL = PatternFill("solid", fgColor="FF1F3864")
ROSA = PatternFill("solid", fgColor="FFFCE4E4")
VERDE = PatternFill("solid", fgColor="FFE2EFDA")
AMAR = PatternFill("solid", fgColor="FFFFF2CC")
GRIS = PatternFill("solid", fgColor="FFF2F2F2")
HFONT = Font(bold=True, size=10, color="FFFFFFFF")
NUM = "#,##0.00"

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Vencimientos"
ws["A1"] = "CIRCA · ZOOM-DIMAX — Cartera por cobrar"
ws["A1"].font = Font(bold=True, size=14)
ws["A2"] = (f"Corte al {fmt_dia(HOY)}-{MES_ES[HOY.month - 1]}-{HOY.year} · {len(cartera)} pedidos "
            f"entregados y sin pago · S/{TOT_CART:,.2f} · vencido S/{TOT_VENC:,.2f} en {len(vencidos)} · "
            f"{len(preventas)} preventas aceptadas sin entregar en su propia hoja")

HDR = ["Supervisor", "Vendedor", "Día de visita", "Pedido", "Bodega", "Celular",
       "Entregado (Perú)", "Vence", "Días", "Situación", "Financiado S/", "Abonado S/",
       "Saldo a cobrar S/", "Mora corre desde", "Cobrable"]
for j, h in enumerate(HDR, 1):
    c = ws.cell(4, j, h)
    c.font, c.fill = HFONT, AZUL
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

r = 5
for x in cartera:
    d = x["dias"]
    sit, fill = ((f"VENCIDO {-d} d", ROSA) if d < 0 else
                 ("vence HOY", AMAR) if d == 0 else (f"por vencer ({d} d)", VERDE))
    valores = [x["supervisor"], x["vendedor"], x["dia_visita"], x["numero"], x["bodega"],
               x["cel"], x["entrega"], x["vence"], d, sit, round(x["mf"], 2),
               round(x["abonado"], 2) if x["abonado"] else None,
               round(x["total"], 2), x["mora_desde"], "SI"]
    for j, v in enumerate(valores, 1):
        c = ws.cell(r, j, v)
        if j in (7, 8, 14):
            c.number_format = "yyyy-mm-dd"
        if j in (11, 12, 13):
            c.number_format = NUM
        if j in (10, 11, 12, 13):
            c.fill = fill
    # Pago parcial: se resalta para que nadie cobre el total completo.
    if x["abonado"]:
        ws.cell(r, 12).font = Font(bold=True, size=10)
    r += 1

fin = r - 1
ws.cell(r, 1, "TOTAL")
ws.cell(r, 11, f"=SUM(K5:K{fin})").number_format = NUM
ws.cell(r, 12, f"=SUM(L5:L{fin})").number_format = NUM
ws.cell(r, 13, f"=SUM(M5:M{fin})").number_format = NUM
for j in range(1, 16):
    ws.cell(r, j).fill = GRIS
    ws.cell(r, j).font = Font(bold=True, size=10)

ws.freeze_panes = "A5"
ws.auto_filter.ref = f"A4:O{fin}"
for col, w in zip("ABCDEFGHIJKLMNO",
                  [30, 26, 14, 10, 42, 15, 16, 12, 7, 18, 13, 12, 17.3, 18, 10]):
    ws.column_dimensions[col].width = w

r += 3
ws.cell(r, 1, "RESUMEN POR SUPERVISOR").font = Font(bold=True, size=11)
r += 1
for j, h in enumerate(["Supervisor", "Pedidos", "Vencidos", "Vencido S/", "Cartera S/"], 1):
    c = ws.cell(r, j, h)
    c.font, c.fill = HFONT, AZUL
for k, v in por_sup.items():
    r += 1
    for j, val in enumerate([k, v[0], v[1], round(v[2], 2), round(v[3], 2)], 1):
        c = ws.cell(r, j, val)
        if j >= 4:
            c.number_format = NUM

r += 3
ws.cell(r, 1, "NOTAS").font = Font(bold=True, size=11)
for nota in [
    "· 'Total a cobrar' NO es el financiado: pasado el día 7 desde la entrega la comisión sube de tramo (1.4% → 3% del día 8 al 14 → 6% del 15 al 30). Calculado con app/services/fees.total_pagar_desde_pedido.",
    "· 'Saldo a cobrar' YA descuenta los abonos (pagos parciales) vigentes de la tabla abonos. Es lo que falta cobrar, no el total original. La columna 'Abonado S/' muestra cuánto entregó la bodega a cuenta.",
    "· 'Mora corre desde' = entrega + 32 días (0.03% diario a partir del día 31 cumplido).",
    "· Los 7 días de plazo corren desde la ENTREGA, no desde la visita.",
    "· Fechas convertidas a hora Perú (UTC-5); la base guarda UTC.",
    "· Vendedor = mapeo activo en bodega_vendedores. Supervisor = vendedores.supervisor (fuente de verdad).",
    "· 'Día de visita' = bodega_vendedores.dia_visita del mapeo activo. Es el día que el vendedor PASA por la bodega, un día antes del despacho — sirve para saber cuándo se puede cobrar en persona. Si dice '— sin ruta', esa bodega no tiene dia_visita en su mapeo y hay que completarlo.",
    "· Solo pedidos entregados y sin pago. Se excluyen bodegas es_test, pedidos TEST-* y pagos en efectivo.",
    "· Los estados cambian durante el día: una bodega puede pagar después de generado este cuadro.",
]:
    r += 1
    ws.cell(r, 1, nota)

# Hoja ventana
ws2 = wb.create_sheet("Por vencer")
dias_ventana = " · ".join(fmt_dia(HOY + timedelta(days=i)) for i in range(3))
ws2["A1"] = f"Vence en los próximos 3 días: {dias_ventana}"
ws2["A1"].font = Font(bold=True, size=12)
ws2["A2"] = f"{len(ventana)} pedidos por S/{TOT_VENT:,.2f}."
for j, h in enumerate(["Vence (día)", "Vence", "Pedido", "Bodega", "Celular", "Vendedor",
                       "Día de visita", "Supervisor", "Financiado S/", "A cobrar S/"], 1):
    c = ws2.cell(4, j, h)
    c.font, c.fill = HFONT, AZUL
r = 5
for x in sorted(ventana, key=lambda y: (y["dias"], -y["total"])):
    for j, v in enumerate([fmt_dia(x["vence"]), x["vence"], x["numero"], x["bodega"], x["cel"],
                           x["vendedor"], x["dia_visita"], x["supervisor"],
                           round(x["mf"], 2), round(x["total"], 2)], 1):
        c = ws2.cell(r, j, v)
        if j == 2:
            c.number_format = "yyyy-mm-dd"
        if j in (9, 10):
            c.number_format = NUM
    r += 1
for col, w in zip("ABCDEFGHIJ", [14, 12, 10, 42, 15, 26, 14, 30, 14, 14]):
    ws2.column_dimensions[col].width = w
ws2.freeze_panes = "A5"

# Hoja preventas
ws3 = wb.create_sheet("Preventas sin entregar")
ws3["A1"] = "Preventas aceptadas — aún SIN ENTREGAR"
ws3["A1"].font = Font(bold=True, size=12)
ws3["A2"] = (f"{len(preventas)} pedidos por S/{sum(x['mf'] for x in preventas):,.2f} de financiamiento "
             f"(S/{sum(x['total'] for x in preventas):,.2f} a cobrar con comisión). NO suman a la cartera: "
             f"el plazo de 7 días arranca el día de la ENTREGA.")
for j, h in enumerate(["Supervisor", "Vendedor", "Día de visita", "Pedido", "Bodega", "Celular",
                       "Estado", "Financiado S/", "A cobrar S/"], 1):
    c = ws3.cell(4, j, h)
    c.font, c.fill = HFONT, AZUL
r = 5
for x in preventas:
    for j, v in enumerate([x["supervisor"], x["vendedor"], x["dia_visita"], x["numero"],
                           x["bodega"], x["cel"], x["estado"],
                           round(x["mf"], 2), round(x["total"], 2)], 1):
        c = ws3.cell(r, j, v)
        if j in (8, 9):
            c.number_format = NUM
    r += 1
for col, w in zip("ABCDEFGHI", [30, 26, 14, 10, 42, 15, 20, 14, 18]):
    ws3.column_dimensions[col].width = w
ws3.freeze_panes = "A5"

xlsx = os.path.join(REPO, f"CIRCA_Vencimientos_{HOY}.xlsx")
wb.save(xlsx)

# ─────────────────────────── correo ───────────────────────────

L = [f"**Asunto:** Cartera Circa al {fmt_fecha(HOY)} — S/{TOT_VENC:,.2f} vencido en {len(vencidos)} bodegas",
     f"**Adjunto:** CIRCA_Vencimientos_{HOY}.xlsx", "", "---", "", "Hola,", ""]

L += ["**Lo que ya venció**", "",
      "| Bodega | Pedido | Abonado | Saldo a cobrar | Vencido hace | Vendedor |",
      "|---|---|---:|---:|---:|---|"]
for x in sorted(vencidos, key=lambda y: -y["total"]):
    d = -x["dias"]
    dd = f"**{d} d**" if d >= 10 else f"{d} d"
    ab = f"S/{x['abonado']:,.2f}" if x["abonado"] else "—"
    L.append(f"| {x['bodega']} | {x['numero']} | {ab} | S/{x['total']:,.2f} | {dd} | "
             f"{solo_apellido(x['vendedor'])} |")
L += [f"| | | | **S/{TOT_VENC:,.2f}** | | |", ""]
_parciales = [x for x in cartera if x["abonado"]]
if _parciales:
    L += [f"> ⚠️ {len(_parciales)} bodega(s) ya abonaron a cuenta. El **saldo** de la "
          f"tabla es lo que falta, no el total original — no cobrar de más.", ""]

if ventana:
    prox = sorted({x["vence"] for x in ventana})
    L += [f"**Lo que vence {' y '.join(fmt_dia(d) for d in prox)}**", "",
          "| Vence | Pedido | Bodega | Celular | A cobrar | Vendedor |",
          "|---|---|---|---|---:|---|"]
    for x in sorted(ventana, key=lambda y: (y["dias"], -y["total"])):
        L.append(f"| {fmt_dia_ab(x['vence'])} | {x['numero']} | {x['bodega']} | {x['cel']} | "
                 f"S/{x['total']:,.2f} | {solo_apellido(x['vendedor'])} |")
    L += [f"| | | | | **S/{TOT_VENT:,.2f}** | |", ""]
else:
    L += ["**No vence nada en los próximos 3 días.**", ""]

concentra = [k for k, v in por_sup.items() if v[2] > 0]
if len(concentra) == 1:
    L.append(f"El 100% del vencido está en un solo equipo")
else:
    L.append("Dónde está el vencido")
L += ["", "| Supervisor | Pedidos | Vencidos | Vencido S/ | Cartera S/ |", "|---|---:|---:|---:|---:|"]
for k, v in por_sup.items():
    neg = "**" if v[2] > 0 else ""
    L.append(f"| {neg}{k}{neg} | {v[0]} | {neg}{v[1]}{neg} | {neg}S/{v[2]:,.2f}{neg} | S/{v[3]:,.2f} |")
L.append("")

foco = concentra[0] if len(concentra) == 1 else None
vs = {k: v for k, v in por_vend.items()
      if foco is None or any(x["vendedor"] == k and x["supervisor"] == foco for x in cartera)}
L += ["| Vendedor | Pedidos | Vencidos | Vencido S/ |", "|---|---:|---:|---:|"]
for k, v in vs.items():
    L.append(f"| {k} | {v[0]} | {v[1]} | S/{v[2]:,.2f} |")

L += ["", "En el excel adjunto el detalle.", "", "Equipo Circa", ""]

md = os.path.join(REPO, f"CIRCA_Correo_Vencimientos_{HOY}.md")
with open(md, "w", encoding="utf-8") as f:
    f.write("\n".join(L))

print(f"✅ {os.path.basename(xlsx)}")
print(f"✅ {os.path.basename(md)}")
print(f"   cartera {len(cartera)} · S/{TOT_CART:,.2f} | vencido {len(vencidos)} · S/{TOT_VENC:,.2f} "
      f"| por vencer {len(ventana)} · S/{TOT_VENT:,.2f} | preventas {len(preventas)}")
